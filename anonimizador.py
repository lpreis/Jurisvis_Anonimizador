from __future__ import annotations

import os
import re
import subprocess
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
import streamlit as st


TOKEN_RE = re.compile(r"\[[^\[\]]+\]")
TOKEN_PARTS_RE = re.compile(r"^\[?([^\[\]\d]+?)[_\s-]*(\d+)\]?$")
TOKEN_TYPE_ALIASES = {
    "PERSON": "PESSOA",
    "PESSOA": "PESSOA",
    "ORGANIZATION": "ORGANIZACAO",
    "ORGANISATION": "ORGANIZACAO",
    "ORGANIZACAO": "ORGANIZACAO",
    "ORGANIZAÇÃO": "ORGANIZACAO",
    "ORG": "ORGANIZACAO",
    "LOCATION": "LOCALIZACAO",
    "LOCALIZACAO": "LOCALIZACAO",
    "LOCALIZAÇÃO": "LOCALIZACAO",
    "ADDRESS": "LOCALIZACAO",
    "MORADA": "LOCALIZACAO",
    "PHONE": "TELEFONE",
    "TELEFONE": "TELEFONE",
    "EMAIL": "EMAIL",
    "NIF": "NIF",
    "NIPC": "NIPC",
    "IBAN": "IBAN",
    "PROCESSO": "PROCESSO",
    "PROCESS": "PROCESSO",
    "CEDULA_PROFISSIONAL": "CEDULA_PROFISSIONAL",
    "CÉDULA_PROFISSIONAL": "CEDULA_PROFISSIONAL",
    "CEDULA": "CEDULA_PROFISSIONAL",
    "FATURA": "FATURA",
    "FACTURA": "FATURA",
    "CONTRATO": "CONTRATO",
    "CARTAO_CIDADAO": "CARTAO_CIDADAO",
    "CARTÃO_CIDADÃO": "CARTAO_CIDADAO",
    "CC": "CARTAO_CIDADAO",
    "CODIGO_POSTAL": "CODIGO_POSTAL",
    "CÓDIGO_POSTAL": "CODIGO_POSTAL",
    "POSTAL_CODE": "CODIGO_POSTAL",
    "CARTAO_CREDITO": "CARTAO_CREDITO",
    "CARTÃO_CRÉDITO": "CARTAO_CREDITO",
    "CREDIT_CARD": "CARTAO_CREDITO",
    "DATA": "DATA",
    "DATE": "DATA",
}
COMPANY_SUFFIX_RE = re.compile(
    r"\b[A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-zÀ-ÿ0-9&.'’-]*(?:\s+(?:de|da|do|das|dos|e|&|[A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-zÀ-ÿ0-9&.'’-]*)){0,5},?\s+(?:Lda\.?|LDA\.?|Limitada|S\.?\s*A\.?|SA)\b\.?",
)
ADDRESS_RE = re.compile(
    r"\b(?:Rua|R\.|Avenida|Av\.|Travessa|Tv\.|Praça|Praca|Praceta|Largo|Estrada|Alameda|Beco|Calçada|Calcada|Rotunda|Urbanização|Urbanizacao|Quinta|Lugar|Caminho)\s+"
    r"[A-ZÁÉÍÓÚÂÊÔÃÕÇ0-9][^,\n]{1,80},?\s+"
    r"(?:n\.?\s*[ºo]\s*)?\d+[A-Za-z]?"
    r"(?:\s*[-/]\s*\d+[A-Za-z]?)?"
    r"(?:\s*,?\s*(?:(?:\d{1,2}\s*\.?\s*(?:º|o|andar))|(?:esq\.?|dto\.?|frente|tr[aá]s)|(?:[A-Z]?\d{1,2}[A-Z])|(?:ap\.?|apt\.?|apartamento|fra[cç][aã]o)\s*[A-Z0-9-]+)){0,3}"
    r"\s*,?\s*\d{4}-\d{3}\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-zÀ-ÿ.'’-]+(?:\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-zÀ-ÿ.'’-]+){0,3}\b",
    flags=re.IGNORECASE,
)


@dataclass(frozen=True)
class EntityMatch:
    start: int
    end: int
    entity_type: str
    text: str
    score: float
    source: str


class ReversibleAnonymizer:
    def __init__(self) -> None:
        self.vault: dict[str, str] = {}
        self.reverse_vault: dict[str, str] = {}
        self.canonical_vault: dict[str, str] = {}
        self.counters: dict[str, int] = {}
        self.presidio_analyzer = self._build_presidio_analyzer()

    def anonymize(self, text: str, language: str = "pt") -> tuple[str, list[EntityMatch]]:
        matches = self.detect(text, language=language)
        anonymized = text
        tokens_by_span: dict[tuple[int, int], str] = {}

        for match in sorted(matches, key=lambda item: item.start):
            tokens_by_span[(match.start, match.end)] = self._token_for(match.entity_type, match.text)

        for match in sorted(matches, key=lambda item: item.start, reverse=True):
            token = tokens_by_span[(match.start, match.end)]
            anonymized = anonymized[: match.start] + token + anonymized[match.end :]

        return anonymized, matches

    def deanonymize(self, text: str) -> str:
        token_aliases = self._token_aliases()

        def replace_token(match: re.Match[str]) -> str:
            token = match.group(0)
            return token_aliases.get(self._normalize_token(token), token)

        return TOKEN_RE.sub(replace_token, text)

    def unresolved_tokens(self, text: str) -> list[str]:
        token_aliases = self._token_aliases()
        unresolved = []
        for token in sorted(set(TOKEN_RE.findall(text))):
            if self._normalize_token(token) not in token_aliases:
                unresolved.append(token)
        return unresolved

    def detect(self, text: str, language: str = "pt") -> list[EntityMatch]:
        matches = self._regex_matches(text)

        if self.presidio_analyzer is not None:
            try:
                presidio_results = self.presidio_analyzer.analyze(
                    text=text,
                    language=language,
                    score_threshold=0.2,
                )
                for result in presidio_results:
                    value = text[result.start : result.end]
                    matches.append(
                        EntityMatch(
                            start=result.start,
                            end=result.end,
                            entity_type=self._normalize_type(result.entity_type),
                            text=value,
                            score=float(result.score),
                            source="presidio",
                        )
                    )
            except Exception as exc:
                st.warning(f"Presidio falhou nesta execução; foi usado o fallback regex. Detalhe: {exc}")

        return self._remove_overlaps(self._refine_matches(matches))

    def reset(self) -> None:
        self.vault.clear()
        self.reverse_vault.clear()
        self.canonical_vault.clear()
        self.counters.clear()

    def _token_for(self, entity_type: str, value: str) -> str:
        existing = self.reverse_vault.get(value)
        if existing:
            return existing

        canonical_key = self._canonical_key(entity_type, value)
        existing = self.canonical_vault.get(canonical_key)
        if existing:
            self.reverse_vault[value] = existing
            return existing

        self.counters[entity_type] = self.counters.get(entity_type, 0) + 1
        if entity_type == "FATURA":
            token = f"[FATURA{self.counters[entity_type]}]"
        else:
            token = f"[{entity_type}_{self.counters[entity_type]}]"
        self.vault[token] = value
        self.reverse_vault[value] = token
        self.canonical_vault[canonical_key] = token
        return token

    def _token_aliases(self) -> dict[str, str]:
        aliases: dict[str, str] = {}
        for token, value in self.vault.items():
            normalized = self._normalize_token(token)
            aliases[normalized] = value

            parts = TOKEN_PARTS_RE.match(token)
            if not parts:
                continue

            token_type, token_number = parts.groups()
            normalized_type = self._normalize_token_type(token_type)
            aliases[f"{normalized_type}_{token_number}"] = value
            aliases[f"{normalized_type}{token_number}"] = value

        return aliases

    @classmethod
    def _normalize_token(cls, token: str) -> str:
        parts = TOKEN_PARTS_RE.match(token.strip())
        if not parts:
            return cls._ascii_upper(token.strip("[]"))

        token_type, token_number = parts.groups()
        normalized_type = cls._normalize_token_type(token_type)
        return f"{normalized_type}_{token_number}"

    @classmethod
    def _normalize_token_type(cls, token_type: str) -> str:
        normalized = cls._ascii_upper(token_type)
        normalized = re.sub(r"[^A-Z0-9]+", "_", normalized).strip("_")
        return TOKEN_TYPE_ALIASES.get(normalized, normalized)

    @staticmethod
    def _ascii_upper(value: str) -> str:
        normalized = unicodedata.normalize("NFKD", value)
        normalized = "".join(char for char in normalized if not unicodedata.combining(char))
        return normalized.upper()

    def _canonical_key(self, entity_type: str, value: str) -> str:
        normalized = unicodedata.normalize("NFKD", value)
        normalized = "".join(char for char in normalized if not unicodedata.combining(char))
        normalized = normalized.casefold().strip()
        normalized = re.sub(r"\s+", " ", normalized)

        if entity_type in {"NIF", "NIPC", "TELEFONE", "CARTAO_CREDITO"}:
            return f"{entity_type}:{re.sub(r'\D', '', normalized)}"

        if entity_type in {"PROCESSO", "FATURA", "CONTRATO"}:
            return f"{entity_type}:{re.sub(r'[^a-z0-9]', '', normalized)}"

        if entity_type == "ORGANIZACAO":
            normalized = re.sub(r"^(?:o|a|os|as)\s+", "", normalized)
            normalized = normalized.replace("&", " e ")
            normalized = re.sub(r"\bs\.?\s*a\.?\b", "sa", normalized)
            normalized = re.sub(r"\bl\.?\s*d\.?\s*a\.?\b", "lda", normalized)
            normalized = re.sub(r"\blimitada\b", "lda", normalized)
            normalized = re.sub(r"[,.;:()\"']", " ", normalized)
            normalized = re.sub(r"\s+", " ", normalized).strip()
            return f"{entity_type}:{normalized}"

        if entity_type in {"EMAIL", "IBAN"}:
            return f"{entity_type}:{re.sub(r'\s+', '', normalized)}"

        normalized = re.sub(r"[,.;:()\"']", " ", normalized)
        normalized = re.sub(r"\s+", " ", normalized).strip()
        return f"{entity_type}:{normalized}"

    def _build_presidio_analyzer(self):
        try:
            from presidio_analyzer import AnalyzerEngine, Pattern, PatternRecognizer
            from presidio_analyzer.nlp_engine import NlpEngineProvider
            import spacy

            if not spacy.util.is_package('pt_core_news_lg'):
                return None

            configuration = {
                "nlp_engine_name": "spacy",
                "models": [{"lang_code": "pt", "model_name": "pt_core_news_lg"}],
            }
            provider = NlpEngineProvider(nlp_configuration=configuration)
            analyzer = AnalyzerEngine(
                nlp_engine=provider.create_engine(),
                supported_languages=["pt"],
            )

            custom_patterns = [
                (
                    "CONTRATO",
                    r"\b(?:Contrato\s+)?[A-Z0-9]{2,}(?:-[A-Z0-9]{2,})+-\d{4}/\d{1,10}\b",
                    0.93,
                    ["contrato", "acordo", "protocolo", "adenda"],
                ),
                (
                    "CARTAO_CIDADAO",
                    r"\b(?:Cart[aã]o\s+de\s+Cidad[aã]o|CC)\s+n\.?\s*[ºo]\s*\d{8}\s*\d\s*[A-Z0-9]{3}\b",
                    0.96,
                    ["cartao de cidadao", "cartão de cidadão", "cc", "identificacao", "identificação"],
                ),
                (
                    "FATURA",
                    r"\b(?:FT|FS|FR|FTR|FAC|NC|ND)\s*\d{4}/\d{1,10}\b",
                    0.94,
                    ["fatura", "factura", "faturas", "facturas", "n.º", "numero", "número"],
                ),
                (
                    "LOCALIZACAO",
                    ADDRESS_RE.pattern,
                    0.99,
                    ["morada", "residencia", "residência", "domicilio", "domicílio", "rua", "avenida"],
                ),
                (
                    "NIPC",
                    r"\bN\.?\s*I\.?\s*P\.?\s*C\.?\s*\d{3}\s?\d{3}\s?\d{3}\b",
                    0.97,
                    ["nipc", "pessoa coletiva", "pessoa colectiva", "contribuinte", "empresa"],
                ),
                (
                    "NIF",
                    r"\b(?:N\.?\s*I\.?\s*F\.?\s*)?[123568]\d{2}\s?\d{3}\s?\d{3}\b",
                    0.72,
                    ["nif", "contribuinte", "numero fiscal", "número fiscal"],
                ),
                (
                    "TELEFONE",
                    r"\b(?:\+351\s?)?(?:9[1236]\d{7}|2\d{8})\b",
                    0.95,
                    ["telefone", "telemovel", "telemóvel", "contacto"],
                ),
                (
                    "ORGANIZACAO",
                    COMPANY_SUFFIX_RE.pattern,
                    0.98,
                    ["sociedade", "empresa", "organizacao", "organização", "lda", "s.a"],
                ),
                (
                    "PROCESSO",
                    r"\b(?:Processo\s+n\.?[ºo]\s*)?\d{1,7}/\d{2}\.[A-Z0-9]{1,6}(?:-[A-Z0-9]{1,6})?\b",
                    0.9,
                    ["processo", "proc", "autos"],
                ),
                (
                    "CEDULA_PROFISSIONAL",
                    r"\bC[eé]dula\s+profissional\s+n\.?[ºo]\s*\d{1,7}[A-Z]?\b",
                    0.9,
                    ["cedula", "cédula", "profissional", "advogado", "advogada"],
                ),
            ]

            for entity, pattern, score, context in custom_patterns:
                analyzer.registry.add_recognizer(
                    PatternRecognizer(
                        supported_entity=entity,
                        patterns=[Pattern(name=f"{entity.lower()}_pattern", regex=pattern, score=score)],
                        supported_language="pt",
                        context=context,
                    )
                )

            return analyzer
        except Exception:
            return None

    def _regex_matches(self, text: str) -> list[EntityMatch]:
        patterns: list[tuple[str, str, float]] = [
            ("EMAIL", r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b", 0.95),
            (
                "CONTRATO",
                r"\b(?:Contrato\s+)?[A-Z0-9]{2,}(?:-[A-Z0-9]{2,})+-\d{4}/\d{1,10}\b",
                0.93,
            ),
            (
                "CARTAO_CIDADAO",
                r"\b(?:Cart[aã]o\s+de\s+Cidad[aã]o|CC)\s+n\.?\s*[ºo]\s*\d{8}\s*\d\s*[A-Z0-9]{3}\b",
                0.96,
            ),
            ("FATURA", r"\b(?:FT|FS|FR|FTR|FAC|NC|ND)\s*\d{4}/\d{1,10}\b", 0.94),
            ("LOCALIZACAO", ADDRESS_RE.pattern, 0.99),
            ("NIPC", r"\bN\.?\s*I\.?\s*P\.?\s*C\.?\s*\d{3}\s?\d{3}\s?\d{3}\b", 0.97),
            (
                "ORGANIZACAO",
                COMPANY_SUFFIX_RE.pattern,
                0.98,
            ),
            ("TELEFONE", r"\b(?:\+351\s?)?(?:9[1236]\d{7}|2\d{8})\b", 0.95),
            ("NIF", r"\b(?:N\.?\s*I\.?\s*F\.?\s*)?[123568]\d{2}\s?\d{3}\s?\d{3}\b", 0.72),
            ("IBAN", r"\bPT50\s?(?:\d{4}\s?){5}\d{1}\b", 0.9),
            (
                "PROCESSO",
                r"\b(?:Processo\s+n\.?[ºo]\s*)?\d{1,7}/\d{2}\.[A-Z0-9]{1,6}(?:-[A-Z0-9]{1,6})?\b",
                0.9,
            ),
            (
                "CEDULA_PROFISSIONAL",
                r"\bC[eé]dula\s+profissional\s+n\.?[ºo]\s*\d{1,7}[A-Z]?\b",
                0.9,
            ),
            ("CODIGO_POSTAL", r"\b\d{4}-\d{3}\b", 0.75),
            ("CARTAO_CREDITO", r"\b(?:\d[ -]*?){13,16}\b", 0.65),
            ("PESSOA", r"\b(?:Dr\.?\s*|Dra\.?\s*)?[A-ZÁÉÍÓÚÂÊÔÃÕÇ][a-zà-ÿ]+(?:\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ][a-zà-ÿ]+)+\b", 0.8),
        ]

        matches: list[EntityMatch] = []
        for entity_type, pattern, score in patterns:
            flags = 0 if entity_type in {"ORGANIZACAO", "PESSOA"} else re.IGNORECASE
            for result in re.finditer(pattern, text, flags=flags):
                value = result.group(0)
                if entity_type == "CARTAO_CREDITO" and not self._looks_like_card(value):
                    continue
                matches.append(
                    EntityMatch(
                        start=result.start(),
                        end=result.end(),
                        entity_type=entity_type,
                        text=value,
                        score=score,
                        source="regex",
                    )
                )

        return matches

    @staticmethod
    def _looks_like_card(value: str) -> bool:
        digits = re.sub(r"\D", "", value)
        if len(digits) < 13:
            return False
        checksum = 0
        parity = len(digits) % 2
        for index, char in enumerate(digits):
            number = int(char)
            if index % 2 == parity:
                number *= 2
                if number > 9:
                    number -= 9
            checksum += number
        return checksum % 10 == 0

    @staticmethod
    def _normalize_type(entity_type: str) -> str:
        aliases = {
            "PERSON": "PESSOA",
            "EMAIL_ADDRESS": "EMAIL",
            "PHONE_NUMBER": "TELEFONE",
            "LOCATION": "LOCALIZACAO",
            "ORGANIZATION": "ORGANIZACAO",
            "IBAN_CODE": "IBAN",
            "DATE_TIME": "DATA",
        }
        normalized = aliases.get(entity_type, entity_type)
        return re.sub(r"[^A-Z0-9_]", "_", normalized.upper())

    @staticmethod
    def _refine_matches(matches: list[EntityMatch]) -> list[EntityMatch]:
        refined: list[EntityMatch] = []

        for match in matches:
            if match.entity_type == "ORGANIZACAO":
                company_matches = list(COMPANY_SUFFIX_RE.finditer(match.text))
                if company_matches:
                    company = company_matches[-1]
                    company_text = company.group(0)
                    article_match = re.match(r"^(?:O|A|Os|As)\s+", company_text)
                    article_offset = article_match.end() if article_match else 0
                    refined.append(
                        EntityMatch(
                            start=match.start + company.start() + article_offset,
                            end=match.start + company.end(),
                            entity_type=match.entity_type,
                            text=company_text[article_offset:],
                            score=max(match.score, 0.98),
                            source=match.source,
                        )
                    )
                    continue

            refined.append(match)

        return refined

    @staticmethod
    def _remove_overlaps(matches: list[EntityMatch]) -> list[EntityMatch]:
        ordered = sorted(matches, key=lambda item: (-item.score, -(item.end - item.start), item.start))
        selected: list[EntityMatch] = []

        for match in ordered:
            has_overlap = any(not (match.end <= saved.start or match.start >= saved.end) for saved in selected)
            if not has_overlap:
                selected.append(match)

        return sorted(selected, key=lambda item: item.start)


def extract_text(uploaded_file) -> str:
    if uploaded_file is None:
        return ""

    suffix = uploaded_file.name.rsplit(".", 1)[-1].lower()
    data = uploaded_file.getvalue()

    if suffix == "txt":
        return data.decode("utf-8", errors="replace")
    if suffix == "docx":
        return extract_docx(data)
    if suffix == "pdf":
        return extract_pdf(data)

    raise ValueError("Formato não suportado. Usa TXT, DOCX ou PDF.")


def extract_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Para ler DOCX instala: pip install python-docx") from exc

    document = Document(BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("Para ler PDF instala: pip install pypdf") from exc

    reader = PdfReader(BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(page.strip() for page in pages if page.strip())


def call_llm(provider: str, prompt: str, model: str, api_key: str, base_url: str) -> str:
    if provider == "OpenAI":
        return call_openai(prompt=prompt, model=model, api_key=api_key)
    if provider == "Ollama":
        return call_ollama(prompt=prompt, model=model, base_url=base_url)
    raise ValueError(f"Fornecedor desconhecido: {provider}")


def call_openai(prompt: str, model: str, api_key: str) -> str:
    try:
        from openai import OpenAI
    except ImportError as exc:
        raise RuntimeError("Para usar OpenAI instala: pip install openai") from exc

    client = OpenAI(api_key=api_key or os.getenv("OPENAI_API_KEY"))
    response = client.responses.create(
        model=model,
        input=prompt,
    )
    return response.output_text


def call_ollama(prompt: str, model: str, base_url: str) -> str:
    try:
        import requests
    except ImportError as exc:
        raise RuntimeError("Para usar Ollama instala: pip install requests") from exc

    url = base_url.rstrip("/") + "/api/generate"
    response = requests.post(
        url,
        json={"model": model, "prompt": prompt, "stream": False},
        timeout=120,
    )
    response.raise_for_status()
    return response.json().get("response", "")


def build_prompt(question: str, text: str, is_anonymized: bool) -> str:
    token_instruction = ""
    if is_anonymized:
        token_instruction = (
            "O texto contem placeholders como [PESSOA_1], [ORGANIZACAO_1] ou [EMAIL_1]. "
            "Mantem todos os placeholders exatamente como estao. "
            "Nao traduzas, nao renomeies e nao inventes placeholders.\n\n"
        )

    return (
        f"{token_instruction}"
        "Responde a pergunta usando apenas a informacao relevante do texto.\n\n"
        f"Pergunta:\n{question.strip()}\n\n"
        f"Texto:\n{text.strip()}"
    )


def initialize_state() -> None:
    if "anonymizer" not in st.session_state:
        st.session_state.anonymizer = ReversibleAnonymizer()
    if "source_text" not in st.session_state:
        st.session_state.source_text = ""
    if "anonymized_text" not in st.session_state:
        st.session_state.anonymized_text = ""
    if "llm_response" not in st.session_state:
        st.session_state.llm_response = ""
    if "deanonymized_response" not in st.session_state:
        st.session_state.deanonymized_response = ""
    if "matches" not in st.session_state:
        st.session_state.matches = []


def get_last_modified_time() -> str:
    try:
        timestamp = os.path.getmtime(__file__)
        return datetime.fromtimestamp(timestamp).strftime("%Y-%m-%d %H:%M:%S")
    except OSError:
        return "desconhecida"


def get_git_version() -> str:
    try:
        sha = subprocess.check_output(
            ["git", "rev-parse", "--short", "HEAD"],
            cwd=os.path.dirname(os.path.abspath(__file__)),
            stderr=subprocess.DEVNULL,
            text=True,
        ).strip()
        return sha
    except Exception:
        return "desconhecida"


def entity_table(matches: list[EntityMatch]) -> list[dict[str, object]]:
    return [
        {
            "tipo": match.entity_type,
            "valor": match.text,
            "score": round(match.score, 2),
            "origem": match.source,
        }
        for match in matches
    ]


def render_sidebar() -> tuple[str, str, str, str]:
    st.sidebar.header("LLM")
    provider = st.sidebar.selectbox("Fornecedor", ["OpenAI", "Ollama"])

    if provider == "OpenAI":
        model = st.sidebar.text_input("Modelo", value="gpt-4.1-mini")
        api_key = st.sidebar.text_input("OPENAI_API_KEY", type="password", value="")
        base_url = ""
        st.sidebar.caption("Também podes definir a variável de ambiente OPENAI_API_KEY.")
    else:
        model = st.sidebar.text_input("Modelo", value="llama3.1")
        base_url = st.sidebar.text_input("URL Ollama", value="http://localhost:11434")
        api_key = ""

    st.sidebar.header("Sessão")
    if st.sidebar.button("Limpar vault e respostas"):
        st.session_state.anonymizer.reset()
        st.session_state.anonymized_text = ""
        st.session_state.llm_response = ""
        st.session_state.deanonymized_response = ""
        st.session_state.matches = []
        st.rerun()

    return provider, model, api_key, base_url


def main() -> None:
    st.set_page_config(page_title="Anonimizador LLM", page_icon="🔐", layout="wide")
    initialize_state()
    provider, model, api_key, base_url = render_sidebar()

    # Debug: show if Presidio is available
    st.write(f"Presidio analyzer available: {st.session_state.anonymizer.presidio_analyzer is not None}")

    st.title("Anonimizador reversível para LLM")
    st.caption(
        f"Última alteração do código: {get_last_modified_time()} | Versão: {get_git_version()}"
    )

    upload_col, text_col = st.columns([0.9, 1.1], gap="large")
    with upload_col:
        uploaded_file = st.file_uploader("Documento TXT, DOCX ou PDF", type=["txt", "docx", "pdf"])
        if uploaded_file is not None and st.button("Extrair texto do documento"):
            try:
                st.session_state.source_text = extract_text(uploaded_file)
            except Exception as exc:
                st.error(str(exc))

        language = st.selectbox("Idioma de deteção", ["pt"], index=0)
        use_anonymized = st.toggle("Enviar texto anonimizado para a LLM", value=True)

    with text_col:
        st.session_state.source_text = st.text_area(
            "Texto base",
            value=st.session_state.source_text,
            height=260,
            placeholder="Escreve ou cola aqui o texto a analisar.",
        )

    actions_col, result_col = st.columns([0.9, 1.1], gap="large")

    with actions_col:
        if st.button("Anonimizar texto", type="primary", use_container_width=True):
            st.session_state.anonymizer.reset()
            anonymized, matches = st.session_state.anonymizer.anonymize(
                st.session_state.source_text,
                language=language,
            )
            st.session_state.anonymized_text = anonymized
            st.session_state.matches = matches
            st.session_state.llm_response = ""
            st.session_state.deanonymized_response = ""

        st.text_area("Texto anonimizado", value=st.session_state.anonymized_text, height=260)

        # Debug: show matches
        if st.session_state.matches:
            st.write(f"Detected {len(st.session_state.matches)} entities:")
            st.dataframe(entity_table(st.session_state.matches), use_container_width=True, hide_index=True)
        else:
            st.info("Ainda não há entidades detetadas nesta sessão.")

    with result_col:
        question = st.text_area(
            "Pergunta para a LLM",
            height=100,
            placeholder="Ex.: resume os factos relevantes, identifica riscos, prepara uma minuta...",
        )

        llm_text = st.session_state.anonymized_text if use_anonymized else st.session_state.source_text
        can_call_llm = bool(question.strip()) and bool(llm_text.strip())

        if st.button("Enviar para a LLM", disabled=not can_call_llm, use_container_width=True):
            try:
                prompt = build_prompt(question=question, text=llm_text, is_anonymized=use_anonymized)
                st.session_state.llm_response = call_llm(
                    provider=provider,
                    prompt=prompt,
                    model=model,
                    api_key=api_key,
                    base_url=base_url,
                )
                st.session_state.deanonymized_response = ""
            except Exception as exc:
                st.error(str(exc))

        st.text_area("Resposta da LLM", height=220, key="llm_response")

        if st.button(
            "Desanonimizar resposta",
            disabled=not bool(st.session_state.llm_response.strip()),
            use_container_width=True,
        ):
            st.session_state.deanonymized_response = st.session_state.anonymizer.deanonymize(
                st.session_state.llm_response
            )
            unresolved_tokens = st.session_state.anonymizer.unresolved_tokens(st.session_state.llm_response)
            if unresolved_tokens:
                st.warning(
                    "Alguns tokens não existem no vault desta sessão: "
                    + ", ".join(unresolved_tokens)
                )

        st.text_area("Resposta desanonimizada", value=st.session_state.deanonymized_response, height=220)

    with st.expander("Vault da sessão"):
        st.caption("Este mapa fica só no estado da sessão da app. Em produção, guarda-o cifrado e com TTL.")
        st.json(st.session_state.anonymizer.vault)

    with st.expander("Verificação rápida de tokens"):
        unknown_tokens = st.session_state.anonymizer.unresolved_tokens(st.session_state.llm_response)
        if unknown_tokens:
            st.warning(f"Tokens sem correspondência no vault: {', '.join(unknown_tokens)}")
        else:
            st.write("Sem tokens desconhecidos na resposta atual.")


if __name__ == "__main__":
    main()
